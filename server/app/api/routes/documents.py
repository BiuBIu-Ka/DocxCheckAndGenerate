
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import FileResponse
from starlette.background import BackgroundTask
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
from app.db.database import get_db
from app.db.models import Document, GjbRule, TermBase
from app.schemas import DocumentCreate, DocumentUpdate, DocumentSchema
from app.services.word_template_parser import word_template_parser
from app.services.template_docx_service import template_docx_service
from typing import List
from pathlib import Path
import shutil
import os
import tempfile
import json
from docx import Document as DocxDocument

router = APIRouter()
STORAGE_ROOT = Path(__file__).resolve().parents[3] / "storage" / "templates"


async def build_template_knowledge_snapshot(db: AsyncSession, doc_type: str) -> tuple[str, str]:
    rules_result = await db.execute(
        select(GjbRule).where(GjbRule.doc_type == doc_type, GjbRule.is_active == True)
    )
    term_result = await db.execute(select(TermBase))

    rules_json = json.dumps(
        [
            {
                "sectionName": rule.section_name,
                "requirementType": rule.requirement_type,
                "description": rule.description,
                "isActive": rule.is_active,
            }
            for rule in rules_result.scalars().all()
        ],
        ensure_ascii=False,
    )
    terms_json = json.dumps(
        [
            {
                "standardName": term.standard_name,
                "aliases": term.aliases,
                "forbiddenTerms": term.forbidden_terms,
                "description": term.description,
            }
            for term in term_result.scalars().all()
        ],
        ensure_ascii=False,
    )
    return rules_json, terms_json


def validate_json_payload(field_name: str, value: str | None) -> None:
    if value is None:
        return
    try:
        parsed = json.loads(value)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail=f"{field_name} 不是合法的 JSON 数据。") from exc

    if field_name in {"structureJson", "rulesJson", "termsJson"} and not isinstance(parsed, list):
        raise HTTPException(status_code=400, detail=f"{field_name} 必须为数组结构。")


def save_template_file(document_id: int, file: UploadFile) -> Path:
    template_dir = STORAGE_ROOT / str(document_id)
    template_dir.mkdir(parents=True, exist_ok=True)
    suffix = Path(file.filename or "").suffix.lower() or ".docx"
    target_path = template_dir / f"source{suffix}"
    with target_path.open("wb") as target:
        shutil.copyfileobj(file.file, target)
    return target_path

@router.get("", response_model=List[DocumentSchema])
async def list_documents(db: AsyncSession = Depends(get_db)):
    result = await db.execute(select(Document).order_by(Document.updated_at.desc()))
    return result.scalars().all()

@router.post("", response_model=DocumentSchema)
async def create_document(doc: DocumentCreate, db: AsyncSession = Depends(get_db)):
    rules_json, terms_json = await build_template_knowledge_snapshot(db, doc.doc_type)
    db_doc = Document(**doc.dict(), rules_json=rules_json, terms_json=terms_json)
    db.add(db_doc)
    await db.commit()
    await db.refresh(db_doc)
    return db_doc

@router.get("/{id}", response_model=DocumentSchema)
async def get_document(id: int, db: AsyncSession = Depends(get_db)):
    db_doc = await db.get(Document, id)
    if not db_doc:
        raise HTTPException(status_code=404, detail="Document not found")
    return db_doc

@router.put("/{id}", response_model=DocumentSchema)
async def update_document(id: int, doc_update: DocumentUpdate, db: AsyncSession = Depends(get_db)):
    db_doc = await db.get(Document, id)
    if not db_doc:
        raise HTTPException(status_code=404, detail="Document not found")

    payload = doc_update.dict(exclude_unset=True)
    validate_json_payload("structureJson", payload.get("structure_json"))
    validate_json_payload("rulesJson", payload.get("rules_json"))
    validate_json_payload("termsJson", payload.get("terms_json"))

    for key, value in payload.items():
        setattr(db_doc, key, value)

    await db.commit()
    await db.refresh(db_doc)
    return db_doc

@router.delete("/{id}")
async def delete_document(id: int, db: AsyncSession = Depends(get_db)):
    db_doc = await db.get(Document, id)
    if not db_doc:
        raise HTTPException(status_code=404, detail="Document not found")
    await db.delete(db_doc)
    await db.commit()
    template_dir = STORAGE_ROOT / str(id)
    if template_dir.exists():
        shutil.rmtree(template_dir, ignore_errors=True)
    return {"success": True}


@router.post("/{id}/template-file", response_model=DocumentSchema)
async def upload_template_file(id: int, file: UploadFile = File(...), db: AsyncSession = Depends(get_db)):
    db_doc = await db.get(Document, id)
    if not db_doc:
        raise HTTPException(status_code=404, detail="Document not found")

    filename = file.filename or ""
    suffix = Path(filename).suffix.lower()
    if suffix not in {".doc", ".docx"}:
        raise HTTPException(status_code=400, detail="模板解析仅支持 .doc 或 .docx 文件，请上传 Word 模板。")

    try:
        saved_path = save_template_file(id, file)
        parsed = word_template_parser.parse_template(str(saved_path), filename)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"模板解析失败，请确认文件是有效的 Word .docx 模板: {exc}") from exc

    db_doc.template_file_name = filename
    db_doc.template_file_path = str(saved_path)
    db_doc.template_html = parsed["html"]
    db_doc.structure_json = json.dumps(parsed["structure"], ensure_ascii=False)
    await db.commit()
    await db.refresh(db_doc)
    return db_doc


@router.get("/{id}/export-docx")
async def export_document_docx(id: int, db: AsyncSession = Depends(get_db)):
    db_doc = await db.get(Document, id)
    if not db_doc:
        raise HTTPException(status_code=404, detail="Document not found")
    if not db_doc.content_json:
        raise HTTPException(status_code=400, detail="当前模板尚未生成文档内容，无法导出。")

    try:
        content_map = json.loads(db_doc.content_json)
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=400, detail="文档内容损坏，无法导出。") from exc

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    temp_path = temp_file.name
    temp_file.close()
    structure = []
    if db_doc.structure_json:
        try:
            structure = json.loads(db_doc.structure_json)
        except json.JSONDecodeError:
            structure = []

    template_path = Path(db_doc.template_file_path) if db_doc.template_file_path else None
    if template_path and template_path.exists() and template_path.suffix.lower() == ".docx":
        template_docx_service.export_with_template(
            str(template_path),
            temp_path,
            content_map,
            structure,
        )
    else:
        # For .doc templates, we cannot use them as a native style base, but we can still 
        # export the markdown formatted content properly into a new .docx.
        export_doc = DocxDocument()
        export_doc.add_heading(db_doc.title, 0)
        export_doc.add_paragraph(f"项目：{db_doc.project_name}")
        export_doc.add_paragraph(f"文档类型：{db_doc.doc_type}")
        if db_doc.generation_prompt:
            export_doc.add_paragraph(f"生成要求：{db_doc.generation_prompt}")

        for title, body in content_map.items():
            heading = export_doc.add_heading(str(title), level=1)
            # Use the new markdown insertion method
            template_docx_service._insert_markdown_after(heading, str(body))
            
        export_doc.save(temp_path)
    download_name = f"{db_doc.title or 'generated-document'}.docx"
    return FileResponse(
        temp_path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=download_name,
        background=BackgroundTask(lambda: Path(temp_path).unlink(missing_ok=True)),
    )


@router.post("/{id}/sync-knowledge", response_model=DocumentSchema)
async def sync_document_knowledge(id: int, db: AsyncSession = Depends(get_db)):
    db_doc = await db.get(Document, id)
    if not db_doc:
        raise HTTPException(status_code=404, detail="Document not found")

    rules_json, terms_json = await build_template_knowledge_snapshot(db, db_doc.doc_type)
    db_doc.rules_json = rules_json
    db_doc.terms_json = terms_json
    await db.commit()
    await db.refresh(db_doc)
    return db_doc

@router.post("/parse-template")
async def parse_template(file: UploadFile = File(...)):
    filename = file.filename or ""
    suffix = Path(filename).suffix.lower()
    if suffix not in {".doc", ".docx"}:
        raise HTTPException(status_code=400, detail="模板解析仅支持 .doc 或 .docx 文件，请上传 Word 模板。")

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    temp_path = temp_file.name
    try:
        with temp_file as buffer:
            shutil.copyfileobj(file.file, buffer)
    except Exception as exc:
        if os.path.exists(temp_path):
            os.remove(temp_path)
        raise HTTPException(status_code=500, detail=f"模板文件保存失败: {exc}") from exc

    try:
        parsed = word_template_parser.parse_template(temp_path, filename)
        return {
            "structure": parsed["structure"],
            "html": parsed["html"],
            "templateFileName": filename,
            "parserKind": parsed.get("parserKind"),
            "fidelity": parsed.get("fidelity"),
        }
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"模板解析失败，请确认文件是有效的 Word .docx 模板: {exc}") from exc
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)
