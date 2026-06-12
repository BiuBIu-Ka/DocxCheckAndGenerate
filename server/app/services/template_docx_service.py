from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


class TemplateDocxService:
    def _normalize(self, text: str) -> str:
        return re.sub(r"[\s\u3000]+", "", (text or "")).lower()

    def _is_heading_text(self, text: str) -> bool:
        text = (text or "").strip()
        if not text:
            return False
        if re.match(r"^\d+(\.\d+)*[\s、.．]+.+", text):
            return True
        return len(text) <= 40 and not text.endswith(("。", "；", ";", ":", "："))

    def _is_heading_paragraph(self, paragraph: Paragraph) -> bool:
        style_name = paragraph.style.name if paragraph.style else ""
        return style_name.startswith("Heading") or self._is_heading_text(paragraph.text)

    def _section_ranges(self, doc: DocxDocument) -> list[dict[str, Any]]:
        sections: list[dict[str, Any]] = []
        paragraphs = list(doc.paragraphs)
        current: dict[str, Any] | None = None
        for index, paragraph in enumerate(paragraphs):
            text = (paragraph.text or "").strip()
            if not text:
                continue
            if self._is_heading_paragraph(paragraph):
                if current is not None:
                    current["end"] = index
                    sections.append(current)
                current = {
                    "title": text,
                    "normalized": self._normalize(text),
                    "start": index,
                    "end": len(paragraphs),
                    "body_style": None,
                }
                continue
            if current is not None and current["body_style"] is None:
                current["body_style"] = paragraph.style.name if paragraph.style else None

        if current is not None:
            sections.append(current)
        return sections

    def _remove_paragraph(self, paragraph: Paragraph) -> None:
        element = paragraph._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

    def _insert_paragraph_after(self, paragraph: Paragraph, text: str, style_name: str | None) -> Paragraph:
        new_p = OxmlElement("w:p")
        paragraph._p.addnext(new_p)
        new_paragraph = Paragraph(new_p, paragraph._parent)
        if style_name:
            try:
                new_paragraph.style = style_name
            except Exception:
                pass
        if text:
            new_paragraph.add_run(text)
        return new_paragraph

    def _replace_section_body(
        self,
        doc: DocxDocument,
        section: dict[str, Any],
        body: str,
    ) -> None:
        paragraphs = list(doc.paragraphs)
        start = section["start"]
        end = section["end"]
        for paragraph in paragraphs[start + 1:end]:
            if self._is_heading_paragraph(paragraph):
                continue
            self._remove_paragraph(paragraph)

        anchor = list(doc.paragraphs)[start]
        style_name = section.get("body_style") or "Normal"
        lines = str(body).splitlines() or [""]
        for line in lines:
            anchor = self._insert_paragraph_after(anchor, line, style_name)

    def _infer_heading_style(self, title: str) -> str:
        if re.match(r"^\d+\.\d+\.\d+", title):
            return "Heading 3"
        if re.match(r"^\d+\.\d+", title):
            return "Heading 2"
        return "Heading 1"

    def export_with_template(
        self,
        template_path: str,
        output_path: str,
        content_map: dict[str, Any],
        structure: list[dict[str, Any]] | None = None,
    ) -> None:
        doc = DocxDocument(template_path)
        sections = self._section_ranges(doc)
        section_map = {section["normalized"]: section for section in sections}

        for title, body in content_map.items():
            section = section_map.get(self._normalize(str(title)))
            if section:
                self._replace_section_body(doc, section, str(body))
                continue

            heading = doc.add_paragraph(style=self._infer_heading_style(str(title)))
            heading.add_run(str(title))
            body_style = "Normal"
            for line in str(body).splitlines() or [""]:
                paragraph = doc.add_paragraph(style=body_style)
                if line:
                    paragraph.add_run(line)

        doc.save(output_path)


template_docx_service = TemplateDocxService()
