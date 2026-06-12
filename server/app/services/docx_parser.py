
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from html import escape
import re
from typing import List, Dict, Any


class DocxParser:
    def _alignment_style(self, paragraph) -> str:
        alignment = paragraph.alignment
        if alignment == WD_ALIGN_PARAGRAPH.CENTER:
            return "text-align:center;"
        if alignment == WD_ALIGN_PARAGRAPH.RIGHT:
            return "text-align:right;"
        if alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
            return "text-align:justify;"
        return ""

    def _render_runs(self, paragraph) -> str:
        fragments: list[str] = []
        if not paragraph.runs:
            return escape(paragraph.text or "")

        for run in paragraph.runs:
            text = escape(run.text or "").replace("\n", "<br />")
            if not text:
                continue
            styles: list[str] = []
            if run.bold:
                styles.append("font-weight:700;")
            if run.italic:
                styles.append("font-style:italic;")
            if run.underline:
                styles.append("text-decoration:underline;")
            if run.font and run.font.size:
                styles.append(f"font-size:{run.font.size.pt:.0f}pt;")
            if run.font and run.font.name:
                styles.append(f"font-family:'{escape(run.font.name)}';")

            if styles:
                fragments.append(f"<span style=\"{''.join(styles)}\">{text}</span>")
            else:
                fragments.append(text)
        return "".join(fragments) or escape(paragraph.text or "")

    def _paragraph_to_html(self, paragraph) -> str:
        text = (paragraph.text or "").strip()
        style_name = paragraph.style.name if paragraph.style else ""
        align_style = self._alignment_style(paragraph)

        if not text:
            return "<p><br /></p>"

        if style_name.startswith("Heading"):
            try:
                level = int(style_name.split(" ")[-1])
            except ValueError:
                level = 1
            level = min(max(level, 1), 6)
            return f"<h{level} style=\"{align_style}\">{self._render_runs(paragraph)}</h{level}>"

        return f"<p style=\"{align_style}\">{self._render_runs(paragraph)}</p>"

    def _table_to_html(self, table) -> str:
        rows_html: list[str] = []
        for row in table.rows:
            cells_html: list[str] = []
            for cell in row.cells:
                paragraphs = [self._paragraph_to_html(paragraph) for paragraph in cell.paragraphs if paragraph.text.strip()]
                cell_content = "".join(paragraphs) or "<p><br /></p>"
                cells_html.append(f"<td>{cell_content}</td>")
            rows_html.append(f"<tr>{''.join(cells_html)}</tr>")
        return "<table border=\"1\" cellspacing=\"0\" cellpadding=\"8\" style=\"width:100%;border-collapse:collapse;\">" + "".join(rows_html) + "</table>"

    def _infer_structure(self, paragraphs: list) -> List[Dict[str, Any]]:
        structure: List[Dict[str, Any]] = []
        current_section = None

        for para in paragraphs:
            text = (para.text or "").strip()
            if not text:
                continue

            style_name = para.style.name if para.style else ""
            is_heading = style_name.startswith("Heading")
            numbered_title = re.match(r"^\d+(\.\d+)*[\s、.．]+.+", text)
            short_title = len(text) <= 40 and not text.endswith(("。", "；", ";", ":", "："))

            if is_heading or numbered_title or short_title:
                level = 1
                if is_heading:
                    try:
                        level = int(style_name.split(" ")[-1])
                    except ValueError:
                        level = 1
                elif numbered_title:
                    level = max(1, text.split()[0].count(".") + 1 if " " in text else text.count(".") + 1)

                section = {
                    "title": text,
                    "level": min(level, 6),
                    "content": "",
                    "children": []
                }
                if level <= 1 or not structure:
                    structure.append(section)
                    current_section = section
                elif current_section:
                    current_section["children"].append(section)
                continue

            if current_section:
                current_section["content"] += text + "\n"

        if structure:
            return structure

        full_text = "\n".join([(para.text or "").strip() for para in paragraphs if (para.text or "").strip()])
        title = "模板正文"
        if full_text:
            title = full_text.splitlines()[0][:40]
        return [{
            "title": title or "模板正文",
            "level": 1,
            "content": full_text,
            "children": []
        }]

    def parse_structure(self, file_path: str) -> List[Dict[str, Any]]:
        doc = Document(file_path)
        return self._infer_structure(doc.paragraphs)

    def parse_template(self, file_path: str) -> Dict[str, Any]:
        doc = Document(file_path)
        structure = self._infer_structure(doc.paragraphs)

        html_blocks: list[str] = [
            "<div class=\"template-doc\" style=\"font-family:'Microsoft YaHei', Arial, sans-serif;line-height:1.7;color:#1f2937;\">"
        ]
        for paragraph in doc.paragraphs:
            html_blocks.append(self._paragraph_to_html(paragraph))
        for table in doc.tables:
            html_blocks.append(self._table_to_html(table))
        html_blocks.append("</div>")

        return {
            "structure": structure,
            "html": "".join(html_blocks),
        }

docx_parser = DocxParser()
