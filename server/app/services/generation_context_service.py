from __future__ import annotations

from pathlib import Path
from typing import Any
import json
import zipfile

from docx import Document as DocxDocument

from app.services.pdf_parser import pdf_parser


TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".rst",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
    ".csv",
    ".log",
    ".py",
    ".js",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".c",
    ".cpp",
    ".h",
    ".hpp",
    ".go",
    ".rs",
    ".sh",
    ".sql",
}


class GenerationContextService:
    def extract_text(self, file_path: str, filename: str) -> str:
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf":
            return pdf_parser.parse_text(file_path)
        if suffix == ".docx":
            return self._extract_docx_text(file_path)
        if suffix == ".zip":
            return self._extract_zip_text(file_path)
        if suffix in TEXT_EXTENSIONS:
            return Path(file_path).read_text(encoding="utf-8", errors="ignore")
        return ""

    def _extract_docx_text(self, file_path: str) -> str:
        doc = DocxDocument(file_path)
        paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
        table_texts: list[str] = []
        for table in doc.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    table_texts.append(" | ".join(values))
        return "\n".join(paragraphs + table_texts)

    def _extract_zip_text(self, file_path: str) -> str:
        snippets: list[str] = []
        with zipfile.ZipFile(file_path) as archive:
            for index, member in enumerate(archive.infolist()):
                if index >= 50:
                    break
                if member.is_dir():
                    continue
                suffix = Path(member.filename).suffix.lower()
                if suffix not in TEXT_EXTENSIONS:
                    continue
                try:
                    raw = archive.read(member)
                except Exception:
                    continue
                text = raw.decode("utf-8", errors="ignore").strip()
                if not text:
                    continue
                snippets.append(f"文件：{member.filename}\n{text[:6000]}")
        return "\n\n".join(snippets)

    def build_context_payload(self, files: list[dict[str, Any]]) -> tuple[str, str]:
        snippets: list[str] = []
        meta: list[dict[str, Any]] = []
        for item in files:
            filename = item["filename"]
            text = (item.get("text") or "").strip()
            excerpt = text[:4000]
            meta.append({
                "filename": filename,
                "chars": len(text),
                "excerpt": excerpt[:800],
            })
            if excerpt:
                snippets.append(f"文件：{filename}\n{excerpt}")
        return "\n\n".join(snippets), json.dumps(meta, ensure_ascii=False)


generation_context_service = GenerationContextService()
